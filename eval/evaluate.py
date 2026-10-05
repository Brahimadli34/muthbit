"""سكربت تقييم مُثبِت ومقارنته بنموذج لغوي عام.

الاستخدام:
  python evaluate.py --xlsx muthbit_test_set.xlsx --api http://localhost:8000 \
      --baseline --repeats 3 --sweep --out reports

يقرأ ملف الاختبار كما هو (ورقتا «مجموعة الاختبار» و«منشورات مركبة»)،
ويُخرج reports/report.md و reports/results.csv.
"""
import argparse
import csv
import importlib.util
import json
import os
import re
import statistics
import sys
import time
from collections import Counter, defaultdict
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl
import requests

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "backend"))


def _load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


# نستورد منطق الحكم والتطبيع نفسيهما المستخدمين في الخلفية
V = _load("verdict", ROOT / "backend/verifier/verdict.py")
AR = _load("arabic", ROOT / "backend/verifier/arabic.py")
LLM = _load("llm_client", ROOT / "backend/verifier/llm_client.py")
LABELS = list(V.LABELS.values())
CONFIRMING_LABELS = {V.LABELS[c] for c in V.CONFIRMING}
ABSTAIN_LABEL = V.LABELS[V.ABSTAIN]

# ---------------------------------------------------------------- قراءة الملف
COLS = {
    "id": "المعرّف", "text": "النص كما يُتداول", "category": "الفئة",
    "critical": "خطأ حرج إن أُكّد؟", "source": "المصدر (الكتاب)", "number": "رقم الحديث/الآية",
    "expected": "المخرج المتوقع من الأداة", "status": "حالة المراجعة",
}


def _sheet_rows(ws):
    headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        yield {h: ("" if v is None else str(v).strip()) for h, v in zip(headers, row)}


def load_test_set(path, approved_only):
    wb = openpyxl.load_workbook(path, read_only=True)
    items = []
    for r in _sheet_rows(wb["مجموعة الاختبار"]):
        rid, text, expected = r.get(COLS["id"], ""), r.get(COLS["text"], ""), r.get(COLS["expected"], "")
        if not text or not rid or not re.match(r"^[A-Za-z]+\d+$", rid):  # يتخطى صفوف الأمثلة
            continue
        if expected not in LABELS:
            print(f"تحذير: {rid} بلا مخرج متوقع صالح، تم تخطيه.")
            continue
        if approved_only and r.get(COLS["status"]) != "معتمد":
            continue
        items.append({"id": rid, "text": text, "category": r.get(COLS["category"], ""),
                      "critical": r.get(COLS["critical"]) == "نعم", "expected": expected,
                      "source": r.get(COLS["source"], ""), "number": r.get(COLS["number"], "")})
    posts = []
    if "منشورات مركبة" in wb.sheetnames:
        by_id = {i["id"]: i for i in items}
        for r in _sheet_rows(wb["منشورات مركبة"]):
            pid, body = r.get("معرّف المنشور", ""), r.get("نص المنشور كاملاً كما يُتداول", "")
            if not body or pid.startswith("مثال"):
                continue
            ids = [x.strip() for x in re.split(r"[،,\s]+", r.get("معرّفات النصوص المنسوبة فيه", "")) if x.strip()]
            posts.append({"id": pid, "text": body, "claims": [by_id[i]["text"] for i in ids if i in by_id]})
    return items, posts


# ---------------------------------------------------------------- الأنظمة
def run_muthbit(api, text, extract=False):
    r = requests.post(f"{api.rstrip('/')}/api/verify", json={"text": text, "extract": extract}, timeout=120)
    r.raise_for_status()
    return r.json()


def muthbit_predict(api, item):
    res = run_muthbit(api, item["text"])["results"][0]
    m = res.get("match") or {}
    return {"label": res["verdict_label"], "source": m.get("source", ""), "number": m.get("number", ""),
            "raw": res.get("top_candidate", {})}


BASELINE_SYSTEM = f"""أنت مساعد متخصص في التحقق من النصوص الشرعية المتداولة.
سيُعطى لك نص متداول. حدد حكمه ومصدره.
أجب بـ JSON فقط: {{"verdict": "...", "source": "...", "number": "...", "grade": "..."}}
حيث verdict واحد حرفياً من: {json.dumps(LABELS, ensure_ascii=False)}
source: اسم الكتاب، و number: رقم الحديث أو الآية، أو اتركهما فارغين."""


def baseline_predict(cfg, item):
    raw = LLM.complete(cfg, BASELINE_SYSTEM, item["text"], max_tokens=400)
    raw = re.sub(r"^```(?:json)?|```$", "", raw.strip()).strip()
    try:
        d = json.loads(raw)
    except json.JSONDecodeError:
        m = re.search(r"\{.*\}", raw, re.S)
        d = json.loads(m.group(0)) if m else {}
    label = d.get("verdict", "")
    return {"label": label if label in LABELS else "غير صالح",
            "source": d.get("source", "") or "", "number": str(d.get("number", "") or "")}


def baseline_configs(args):
    """النماذج المرجعية: من لوحة الإدارة (use_for_baseline) أو من سطر الأوامر."""
    if args.baseline_model:
        key = os.environ.get(args.baseline_key_env, "")
        return [LLM.ProviderConfig(name=args.baseline_model, api_style=args.baseline_style,
                                   model=args.baseline_model, api_key=key, base_url=args.baseline_base_url)]
    os.environ.setdefault("DJANGO_SETTINGS_MODULE", "muthbit.settings")
    import django
    django.setup()
    from verifier.models import AIProvider
    cfgs = [p.to_config() for p in AIProvider.objects.filter(is_active=True, use_for_baseline=True)]
    return [LLM.ProviderConfig(**c.__dict__) for c in cfgs if c.api_key]


# ---------------------------------------------------------------- المقاييس
def _digits(s):
    return re.sub(r"\D", "", str(s))


def source_correct(item, pred):
    if not item["source"] or not pred["source"]:
        return False
    a, b = AR.normalize(item["source"]), AR.normalize(pred["source"])
    return (a in b or b in a) and _digits(item["number"]) == _digits(pred["number"]) != ""


def metrics(items, preds):
    n = len(items)
    correct = sum(p["label"] == i["expected"] for i, p in zip(items, preds))
    crit = [(i, p) for i, p in zip(items, preds) if i["critical"]]
    crit_err = sum(p["label"] in CONFIRMING_LABELS for _, p in crit)
    exp_abs = [p for i, p in zip(items, preds) if i["expected"] == ABSTAIN_LABEL]
    pred_abs = [(i, p) for i, p in zip(items, preds) if p["label"] == ABSTAIN_LABEL]
    with_src = [(i, p) for i, p in zip(items, preds) if i["source"] and i["number"]]
    no_src = [(i, p) for i, p in zip(items, preds) if not i["source"]]
    fabricated = sum(bool(p["source"] and p["number"]) for _, p in no_src)
    per_cat = defaultdict(lambda: [0, 0])
    for i, p in zip(items, preds):
        per_cat[i["category"] or "بلا فئة"][1] += 1
        per_cat[i["category"] or "بلا فئة"][0] += p["label"] == i["expected"]
    return {
        "n": n,
        "accuracy": correct / n if n else 0,
        "critical_errors": crit_err, "critical_total": len(crit),
        "abstain_recall": sum(p["label"] == ABSTAIN_LABEL for p in exp_abs) / len(exp_abs) if exp_abs else None,
        "abstain_precision": sum(i["expected"] == ABSTAIN_LABEL for i, _ in pred_abs) / len(pred_abs) if pred_abs else None,
        "attribution_accuracy": sum(source_correct(i, p) for i, p in with_src) / len(with_src) if with_src else None,
        "attribution_total": len(with_src),
        "fabricated_citations": fabricated, "no_source_total": len(no_src),
        "per_category": {k: v[0] / v[1] for k, v in per_cat.items()},
    }


def consistency(runs):
    """نسبة النصوص التي أعطت المخرج نفسه في كل التكرارات."""
    if len(runs) < 2:
        return None
    same = sum(len({r[k]["label"] for r in runs}) == 1 for k in range(len(runs[0])))
    return same / len(runs[0])


def sweep(items, raws):
    """يبحث عن أفضل عتبتين: أقل أخطاء حرجة أولاً، ثم أعلى دقة."""
    best = None
    grid = [x / 100 for x in range(30, 100, 5)]
    for match in grid:
        for exact in grid:
            if exact < match:
                continue
            t = V.Thresholds(exact=exact, match=match)
            preds = [{"label": V.LABELS[V.decide(r.get("score", 0), r.get("grade_class"),
                                                  r.get("has_other_grades", False), t)],
                      "source": "", "number": ""} for r in raws]
            m = metrics(items, preds)
            key = (-m["critical_errors"], m["accuracy"])
            if best is None or key > best[0]:
                best = (key, match, exact, m)
    return best


def evaluate_posts(api, posts):
    """استخراج النصوص من المنشورات المركبة: هل استُخرج كل نص متوقع؟"""
    found = total = exact_count = 0
    for p in posts:
        res = run_muthbit(api, p["text"], extract=True)
        got = [AR.normalize(r["claim"]) for r in res["results"]]
        exact_count += len(got) == len(p["claims"])
        for c in p["claims"]:
            total += 1
            cn = AR.normalize(c)
            found += any(SequenceMatcher(None, cn, g).ratio() >= 0.6 or cn in g or g in cn for g in got)
    return {"posts": len(posts), "claim_recall": found / total if total else None,
            "exact_count_rate": exact_count / len(posts) if posts else None}


# ---------------------------------------------------------------- التقرير
def pct(x):
    return "—" if x is None else f"{x * 100:.1f}%"


def write_report(out, items, systems, sweep_result, posts_metrics, args):
    out.mkdir(parents=True, exist_ok=True)
    with (out / "results.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        head = ["id", "category", "critical", "expected"]
        for name in systems:
            head += [f"{name}_label", f"{name}_source", f"{name}_number", f"{name}_correct"]
        w.writerow(head)
        for k, it in enumerate(items):
            row = [it["id"], it["category"], it["critical"], it["expected"]]
            for s in systems.values():
                p = s["runs"][0][k]
                row += [p["label"], p["source"], p["number"], p["label"] == it["expected"]]
            w.writerow(row)

    L = [f"# تقرير تقييم مُثبِت", "", f"- تاريخ التشغيل: {time.strftime('%Y-%m-%d %H:%M')}",
         f"- عدد النصوص: {len(items)}", f"- عدد التكرارات: {args.repeats}", ""]
    L += ["## المقارنة الرئيسية", "", "| المقياس | " + " | ".join(systems) + " |",
          "|---|" + "---|" * len(systems)]
    rows = [
        ("مطابقة المخرج المتوقع", lambda m: pct(m["accuracy"])),
        ("أخطاء حرجة (تأكيد ما لا يصح)", lambda m: f"{m['critical_errors']} من {m['critical_total']}"),
        ("إسناد صحيح (المصدر والرقم)", lambda m: f"{pct(m['attribution_accuracy'])} من {m['attribution_total']}"),
        ("مصادر مُختلقة لنصوص بلا أصل", lambda m: f"{m['fabricated_citations']} من {m['no_source_total']}"),
        ("استدعاء الامتناع", lambda m: pct(m["abstain_recall"])),
        ("دقة الامتناع", lambda m: pct(m["abstain_precision"])),
    ]
    for title, fn in rows:
        L.append(f"| {title} | " + " | ".join(fn(s["metrics"]) for s in systems.values()) + " |")
    L.append("| ثبات المخرج عبر التكرارات | " + " | ".join(pct(s["consistency"]) for s in systems.values()) + " |")
    L += ["", "## الدقة حسب الفئة", "", "| الفئة | " + " | ".join(systems) + " |", "|---|" + "---|" * len(systems)]
    cats = sorted({c for s in systems.values() for c in s["metrics"]["per_category"]})
    for c in cats:
        L.append(f"| {c} | " + " | ".join(pct(s["metrics"]["per_category"].get(c)) for s in systems.values()) + " |")
    if posts_metrics:
        L += ["", "## استخراج النصوص من المنشورات المركبة", "",
              f"- عدد المنشورات: {posts_metrics['posts']}",
              f"- نسبة النصوص المتوقعة التي استُخرجت: {pct(posts_metrics['claim_recall'])}",
              f"- منشورات استُخرج منها العدد الصحيح تماماً: {pct(posts_metrics['exact_count_rate'])}"]
    if sweep_result:
        _, match, exact, m = sweep_result
        L += ["", "## ضبط العتبات", "",
              f"أفضل عتبتين على هذه المجموعة: THRESHOLD_MATCH={match} و THRESHOLD_EXACT={exact}",
              f"(أخطاء حرجة: {m['critical_errors']}، مطابقة: {pct(m['accuracy'])}).",
              "", "تنبيه: العتبات المضبوطة على مجموعة الاختبار نفسها تبالغ في تقدير الأداء. "
              "للنتيجة المنشورة: اضبط العتبات على نصف المجموعة وقِس على النصف الآخر (--split)."]
    L += ["", "## حدود هذا التقييم", "",
          "- النتائج تخص هذه المجموعة ومصادر الفهرس الحالية، ولا تُعمَّم على كل النصوص المتداولة.",
          "- المخرج المتوقع لكل نص حدده المراجع الشرعي للمشروع."]
    (out / "report.md").write_text("\n".join(L), encoding="utf-8")


# ---------------------------------------------------------------- التشغيل
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--api", default=os.environ.get("MUTHBIT_API", "http://localhost:8000"))
    ap.add_argument("--baseline", action="store_true", help="مقارنة بنموذج لغوي عام")
    ap.add_argument("--baseline-model", help="بدون هذا الخيار تُقرأ النماذج المرجعية من لوحة الإدارة")
    ap.add_argument("--baseline-style", default="openai", choices=["openai", "anthropic"])
    ap.add_argument("--baseline-base-url", default="")
    ap.add_argument("--baseline-key-env", default="BASELINE_API_KEY")
    ap.add_argument("--repeats", type=int, default=1)
    ap.add_argument("--sweep", action="store_true", help="ضبط العتبات")
    ap.add_argument("--split", action="store_true", help="ضبط على النصف الأول والقياس على الثاني")
    ap.add_argument("--approved-only", action="store_true", help="النصوص «المعتمدة» فقط")
    ap.add_argument("--out", default="reports")
    args = ap.parse_args()

    items, posts = load_test_set(args.xlsx, args.approved_only)
    if not items:
        sys.exit("لا توجد نصوص صالحة في ملف الاختبار.")
    print(f"نصوص: {len(items)}، منشورات مركبة: {len(posts)}")

    predictors = {"مُثبِت": lambda it: muthbit_predict(args.api, it)}
    if args.baseline:
        cfgs = baseline_configs(args)
        if not cfgs:
            sys.exit("لا يوجد نموذج مرجعي: فعّل «يُستعمل نموذجاً مرجعياً» لنموذج في لوحة الإدارة أو استعمل --baseline-model.")
        for cfg in cfgs:
            predictors[f"نموذج عام ({cfg.model})"] = (lambda c: (lambda it: baseline_predict(c, it)))(cfg)

    systems = {}
    for name, fn in predictors.items():
        runs = []
        for rep in range(args.repeats):
            preds = []
            for k, it in enumerate(items, 1):
                try:
                    preds.append(fn(it))
                except Exception as exc:
                    print(f"  خطأ في {it['id']}: {exc}")
                    preds.append({"label": "خطأ تشغيل", "source": "", "number": "", "raw": {}})
                print(f"\r{name} تكرار {rep + 1}: {k}/{len(items)}", end="", flush=True)
            print()
            runs.append(preds)
        systems[name] = {"runs": runs, "metrics": metrics(items, runs[0]), "consistency": consistency(runs)}

    sweep_result = None
    if args.sweep:
        raws = [p.get("raw", {}) for p in systems["مُثبِت"]["runs"][0]]
        if args.split:
            half = len(items) // 2
            best = sweep(items[:half], raws[:half])
            t = V.Thresholds(exact=best[2], match=best[1])
            held = [{"label": V.LABELS[V.decide(r.get("score", 0), r.get("grade_class"),
                                                 r.get("has_other_grades", False), t)],
                     "source": "", "number": ""} for r in raws[half:]]
            sweep_result = (best[0], best[1], best[2], metrics(items[half:], held))
        else:
            sweep_result = sweep(items, raws)

    posts_metrics = evaluate_posts(args.api, posts) if posts else None
    write_report(Path(args.out), items, systems, sweep_result, posts_metrics, args)
    print(f"تم. التقرير في {Path(args.out) / 'report.md'}")


if __name__ == "__main__":
    main()
